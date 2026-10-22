#!/usr/bin/env python3
"""Export CoTZ configuration to XLSX and apply validated, three-way edits.

Lua is tokenized, never executed. Only unambiguous literal configuration values
are editable. Unknown syntax and executable expressions remain untouched.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
import hashlib
import io
import json
import math
from pathlib import Path
import re
import subprocess
import sys
import tarfile
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Protection

SCHEMA = 2
ROOT = Path(__file__).resolve().parents[2]
ITEM_FIELDS = set("price width height weight flatweight baseweight varweight BaseWeight WeightPerLevel quantity restore restoreduration ammoAmount bulletweight unloadedweight br fbr ar far sr fsr pr fpr radProt flatRadProt rarity JumpPower carryinc exteriorSlots interiorSlots extraSlots hunger thirst alcohol psyheal psyblock radrem WeightPerHunger WeightPerThirst pricepertier upgradeCost meattier fueltier knifetier cookertier maxFuel maxStack".split())
WEAPON_FIELDS = set("Damage FireDelay Primary.ClipSize Primary.DefaultClip Primary.Automatic DurabilityDamageChance WearDamage HipSpread AimSpread Recoil ClumpSpread VelocitySensitivity Shots".split())
MUTANT_FIELDS = set("hp hpvar BR FBR FBRAP MinRangeDist MaxRangeDist".split())
ANOMALY_FIELDS = set("BalanceDamage BalanceCooldown BalanceRadius".split())
POOLS = {"commonArtifacts", "rareArtifacts", "veryRareArtifacts"}
FIELD_FIELDS = {"interval"} | POOLS
READ_FIELDS = {"name", "PrintName", "base", "Base", "category", "entityname"}
INTEGER_FIELDS = {"width", "height", "quantity", "ammoAmount", "Primary.ClipSize", "Primary.DefaultClip", "rarity", "exteriorSlots", "interiorSlots", "extraSlots", "maxStack", "meattier", "fueltier", "knifetier", "cookertier"}
TAB_FIELDS = {
    "Items": ITEM_FIELDS, "Item bases": ITEM_FIELDS, "Weapons": WEAPON_FIELDS,
    "Mutants": MUTANT_FIELDS, "Anomalies": ANOMALY_FIELDS,
    "Anomaly fields": FIELD_FIELDS,
}
IDENTITY = ["ID", "Source", "Name", "Base"]


class BalanceError(ValueError):
    pass


@dataclass
class Token:
    value: str
    start: int
    end: int
    kind: str = "symbol"


def tokenize(source):
    """Skip Lua/GLua comments and retain exact source offsets, including BOMs."""
    result = []
    i = 0
    while i < len(source):
        c = source[i]
        if c.isspace() or c == "\ufeff":
            i += 1
            continue
        if source.startswith(("--", "//"), i):
            long = re.match(r"\[(=*)\[", source[i + 2:]) if source.startswith("--", i) else None
            if long:
                close = "]" + long[1] + "]"
                end = source.find(close, i + 2 + len(long[0]))
                if end < 0:
                    raise BalanceError("Unterminated long comment")
                i = end + len(close)
            else:
                end = source.find("\n", i)
                i = len(source) if end < 0 else end
            continue
        if source.startswith("/*", i):
            end = source.find("*/", i + 2)
            if end < 0:
                raise BalanceError("Unterminated block comment")
            i = end + 2
            continue
        start = i
        if c in "\"'":
            i += 1
            while i < len(source):
                if source[i] == "\\":
                    i += 2
                elif source[i] == c:
                    i += 1
                    break
                else:
                    i += 1
            else:
                raise BalanceError("Unterminated string")
            result.append(Token(source[start:i], start, i, "string"))
            continue
        long = re.match(r"\[(=*)\[", source[i:])
        if long:
            close = "]" + long[1] + "]"
            end = source.find(close, i + len(long[0]))
            if end < 0:
                raise BalanceError("Unterminated long string")
            i = end + len(close)
            result.append(Token(source[start:i], start, i, "longstring"))
            continue
        match = re.match(r"[A-Za-z_][A-Za-z_0-9]*", source[i:])
        kind = "identifier"
        if not match:
            match = re.match(r"(?:0[xX][0-9a-fA-F]+|(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?)", source[i:])
            kind = "number"
        if match:
            i += len(match[0])
            result.append(Token(match[0], start, i, kind))
        else:
            op = next((v for v in ("...", "..", "==", "~=", "!=", "<=", ">=", "&&", "||") if source.startswith(v, i)), c)
            i += len(op)
            result.append(Token(op, start, i))
    return result


def lua_string(raw):
    # Configuration names/classes normally use plain ASCII. Do not guess how an
    # unfamiliar Lua escape is interpreted; preserve it verbatim for display.
    return re.sub(r"\\([\\\"'])", r"\1", raw[1:-1])


def literal(tokens, i, pools=False):
    if i >= len(tokens):
        raise BalanceError("Missing value")
    start = i
    t = tokens[i]
    if pools and t.value == "{":
        values = []
        i += 1
        while i < len(tokens) and tokens[i].value != "}":
            if tokens[i].kind != "string" or "\\" in tokens[i].value:
                raise BalanceError("Artifact pools must contain plain string literals")
            values.append(lua_string(tokens[i].value))
            i += 1
            if i < len(tokens) and tokens[i].value in {",", ";"}:
                i += 1
            elif i < len(tokens) and tokens[i].value != "}":
                raise BalanceError("Unsupported artifact pool")
        if i >= len(tokens):
            raise BalanceError("Unterminated artifact pool")
        return values, i + 1, tokens[start].start, tokens[i].end
    sign = 1
    if t.value == "-":
        sign = -1
        i += 1
        t = tokens[i]
    if t.kind == "number":
        value = sign * (int(t.value, 16) if t.value.lower().startswith("0x") else float(t.value))
        if not math.isfinite(value):
            raise BalanceError("Non-finite number")
        return value, i + 1, tokens[start].start, t.end
    if sign == 1 and t.value in {"true", "false"}:
        return t.value == "true", i + 1, t.start, t.end
    if sign == 1 and t.kind == "string":
        return lua_string(t.value), i + 1, t.start, t.end
    raise BalanceError("Expression is not a supported literal")


def extract(source, prefix, allowed, field_table=False):
    """Find direct assignments outside functions/control blocks and nested tables.

    Return all field names, including unsupported assignments, so a later dynamic
    assignment cannot accidentally leave an earlier literal editable.
    """
    tokens = tokenize(source)
    depth = 0
    blocks = []
    pending_do = 0
    assignments = defaultdict(list)
    parent_writes = []
    for i, t in enumerate(tokens):
        value = t.value
        if not blocks and depth == 0:
            j = i
            if not field_table and value == prefix and i + 1 < len(tokens) and tokens[i + 1].value == "=":
                parent_writes.append(("", t.start))
            if field_table:
                name = value
                j += 1
            else:
                if value != prefix or i + 2 >= len(tokens) or tokens[i + 1].value != ".":
                    name = None
                else:
                    j = i + 2
                    parts = [tokens[j].value]
                    j += 1
                    while j + 1 < len(tokens) and tokens[j].value == ".":
                        parts.append(tokens[j + 1].value)
                        j += 2
                    name = ".".join(parts)
            if name is not None and j < len(tokens) and tokens[j].value == "=":
                if any(field.startswith(name + ".") for field in allowed):
                    parent_writes.append((name, t.start))
            if name in allowed and j < len(tokens) and tokens[j].value == "=":
                try:
                    val, end, a, b = literal(tokens, j + 1, name in POOLS)
                    # Require a statement/table-field boundary, not `10 * scale`.
                    if end < len(tokens):
                        nxt = tokens[end]
                        if field_table:
                            valid_end = nxt.value in {",", ";"}
                        else:
                            valid_end = nxt.value == ";" or ("\n" in source[b:nxt.start] and nxt.value not in {"+", "-", "*", "/", "^", "%", "..", "and", "or", "(", "[", ".", ":"})
                        if not valid_end:
                            raise BalanceError("Value is part of an expression")
                    assignments[name].append({"value": val, "start": a, "end": b})
                except (BalanceError, IndexError):
                    assignments[name].append(None)
        if t.kind != "identifier" and t.kind != "symbol":
            continue
        if value in {"(", "{", "["}:
            depth += 1
        elif value in {")", "}", "]"}:
            depth -= 1
        elif value in {"function", "if", "for", "while", "repeat"}:
            blocks.append(value)
            if value in {"for", "while"}:
                pending_do += 1
        elif value == "do":
            if pending_do:
                pending_do -= 1
            else:
                blocks.append(value)
        elif value in {"end", "until"} and blocks:
            blocks.pop()
    if depth != 0 or blocks or pending_do:
        raise BalanceError("Unbalanced delimiters/control scopes; file is read-only")
    fields = {k: v[0] for k, v in assignments.items() if len(v) == 1 and v[0] is not None}
    issues = [f"{k}: repeated or non-literal assignment (read-only)" for k, v in assignments.items() if len(v) != 1 or v[0] is None]
    for name in list(fields):
        if any((not parent or name.startswith(parent + ".")) and fields[name]["start"] < offset for parent, offset in parent_writes):
            issues.append(f"{name}: containing table overwritten later (read-only)")
            del fields[name]
    return fields, issues


def git(root, *args):
    return subprocess.check_output(["git", "-C", str(root), *args], stderr=subprocess.PIPE).decode().strip()


def item_files(root):
    roots = [root / "gamemodes/ix_cotz/schema/items"]
    roots += sorted((root / "gamemodes/ix_cotz/plugins").glob("*/items"))
    for folder in roots:
        # Match Helix LoadFromDir: root files, then exactly one category level.
        yield from sorted(folder.glob("sh_*.lua"))
        yield from sorted(folder.glob("*/sh_*.lua"))


def scan(root):
    entities = []
    warnings = []

    def add(path, tab, ident, prefix, allowed):
        relative = path.relative_to(root).as_posix()
        source = path.read_bytes().decode("utf-8")
        try:
            fields, issues = extract(source, prefix, allowed | READ_FIELDS)
        except BalanceError as exc:
            warnings.append([relative, str(exc)])
            return
        for issue in issues:
            warnings.append([relative, issue])
        entity = {"tab": tab, "id": ident, "source": relative,
                  "name": fields.get("name", fields.get("PrintName", {})).get("value", ident),
                  "base": fields.get("base", fields.get("Base", {})).get("value", ""),
                  "category": fields.get("category", {}).get("value", "") if prefix == "ITEM" else "",
                  "fields": {k: v for k, v in fields.items() if k in allowed}}
        if not isinstance(entity["category"], str):
            entity["category"] = ""
        if prefix == "ITEM" and not entity["base"] and path.parent.name != "items":
            entity["base"] = "base_" + path.parent.name if path.parent.name != "base" else ""
        # Textual metadata is not a configuration edit target in this version.
        for field in list(entity["fields"]):
            val = entity["fields"][field]["value"]
            if not isinstance(val, (int, float, bool)):
                warnings.append([relative, f"{field}: expected numeric/boolean literal (read-only)"])
                del entity["fields"][field]
        entities.append(entity)
        if tab == "Mutants" and "TEMP_MeleeTable.damage" in source:
            warnings.append([relative, "Attack damage/timing is embedded in initialization functions; not editable yet."])
        if tab == "Anomalies" and re.search(r"TakeDamage\(|SetDamage\(|timer.Create\(", source):
            warnings.append([relative, "Only explicit Balance* attributes are editable; other damage/timers remain in behavior code."])

    for path in item_files(root):
        base = path.parent.name == "base"
        add(path, "Item bases" if base else "Items", ("base_" if base else "") + path.stem[3:], "ITEM", ITEM_FIELDS)
    for path in sorted((root / "lua/weapons").glob("*/shared.lua")):
        add(path, "Weapons", path.parent.name, "SWEP", WEAPON_FIELDS)
    for pattern, tab, allowed in [("npc_mutant_*", "Mutants", MUTANT_FIELDS), ("anom_*", "Anomalies", ANOMALY_FIELDS)]:
        for folder in sorted((root / "lua/entities").glob(pattern)):
            for name in ("shared.lua", "init.lua"):
                path = folder / name
                if path.is_file():
                    add(path, tab, folder.name + "/" + name, "ENT", allowed)
    path = root / "gamemodes/ix_cotz/plugins/anomalycontroller/sh_anomalydefs.lua"
    if path.is_file():
        source = path.read_bytes().decode("utf-8")
        tokens = tokenize(source)
        for i in range(len(tokens) - 7):
            if [t.value for t in tokens[i:i + 4]] != ["PLUGIN", ".", "anomalydefs", "["]:
                continue
            if tokens[i + 4].kind != "number" or [t.value for t in tokens[i + 5:i + 8]] != ["]", "=", "{"]:
                continue
            depth = 1
            j = i + 8
            while j < len(tokens) and depth:
                if tokens[j].value == "{":
                    depth += 1
                elif tokens[j].value == "}":
                    depth -= 1
                j += 1
            if depth:
                raise BalanceError("Unterminated anomaly definition")
            start = tokens[i + 7].end
            fields, issues = extract(source[start:tokens[j - 1].start], "", FIELD_FIELDS | READ_FIELDS, True)
            for val in fields.values():
                val["start"] += start
                val["end"] += start
            warnings.extend([[path.relative_to(root).as_posix(), issue] for issue in issues])
            entities.append({"tab": "Anomaly fields", "id": tokens[i + 4].value,
                             "source": path.relative_to(root).as_posix(),
                             "name": fields.get("name", {}).get("value", ""),
                             "base": fields.get("entityname", {}).get("value", ""),
                             "category": "",
                             "fields": {k: v for k, v in fields.items() if k in FIELD_FIELDS}})
    counts = Counter((e["tab"], e["id"]) for e in entities)
    bases = {e["id"]: e for e in entities if e["tab"] == "Item bases" and counts[e["tab"], e["id"]] == 1}
    for entity in entities:
        if entity["tab"] != "Items":
            continue
        category = entity["category"]
        base = entity["base"]
        visited = set()
        while not category and base in bases and base not in visited:
            visited.add(base)
            category = bases[base]["category"]
            base = bases[base]["base"]
        entity["category"] = category or (Path(entity["source"]).parent.name if Path(entity["source"]).parent.name != "items" else "Uncategorized")
    for entity in entities:
        if counts[entity["tab"], entity["id"]] > 1:
            warnings.append([entity["source"], f"Duplicate ID {entity['id']}; all matching definitions are read-only."])
            entity["fields"] = {}
    return entities, warnings


def display(value):
    return json.dumps(value, ensure_ascii=False) if isinstance(value, list) else value


def digest(data):
    return hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def append_safe(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # Names starting with '=' must remain text.


def workbook_layout(entities):
    """Derive sheets and editable columns from the immutable export baseline."""
    layout = {}
    categories = sorted({e["category"] for e in entities if e["tab"] == "Items"})
    used = {name.casefold() for name in [*TAB_FIELDS, "Metadata", "Coverage", "Column guide"]}
    for category in categories:
        raw = "Items - " + category
        name = re.sub(r"[\\/*?:\[\]]", "_", raw).strip("'")
        if len(name) > 31 or name != raw or name.casefold() in used:
            name = name[:22] + "-" + hashlib.sha256(category.encode()).hexdigest()[:8]
        used.add(name.casefold())
        rows = [e for e in entities if e["tab"] == "Items" and e["category"] == category]
        columns = sorted({field for e in rows for field in e["fields"]})
        layout[name] = (rows, columns)
    for tab, allowed in TAB_FIELDS.items():
        if tab != "Items":
            layout[tab] = ([e for e in entities if e["tab"] == tab], sorted(allowed))
    return layout


def export(root, output, preview=False):
    entities, warnings = scan(root)
    commit = git(root, "rev-parse", "HEAD")
    manifest = {"schema": SCHEMA, "commit": commit, "entities": entities}
    if not preview:
        trusted_baseline(root, manifest)
    output.mkdir(parents=True, exist_ok=True)
    (output / "baseline.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")
    book = Workbook()
    book.remove(book.active)
    meta = book.create_sheet("Metadata")
    for row in [["Schema", SCHEMA], ["Commit", commit], ["Baseline SHA256", digest(manifest)],
                ["Instructions", "Edit blue cells only. IDs and source paths must remain unchanged. Import requires baseline.json."],
                ["Missing values", "Blank means absent/unsupported, not zero. Adding or clearing fields is not supported."],
                ["Inheritance", "Explicit source values only. Edit defaults on Item bases; inherited/runtime values are not expanded."],
                ["Scope", "Definitions discovered using schema/plugin item paths and entity/weapon directories; runtime availability is not verified."]]:
        append_safe(meta, row)
    for tab, (rows, columns) in workbook_layout(entities).items():
        sheet = book.create_sheet(tab)
        append_safe(sheet, IDENTITY + columns)
        for entity in rows:
            append_safe(sheet, [entity["id"], entity["source"], entity["name"], entity["base"]] +
                        [display(entity["fields"][f]["value"]) if f in entity["fields"] else None for f in columns])
            for index, field in enumerate(columns, len(IDENTITY) + 1):
                if field in entity["fields"]:
                    cell = sheet.cell(sheet.max_row, index)
                    cell.fill = PatternFill("solid", fgColor="DDEEFF")
                    cell.protection = Protection(locked=False)
        sheet.freeze_panes = "E2"
        sheet.auto_filter.ref = sheet.dimensions
    coverage = book.create_sheet("Coverage")
    append_safe(coverage, ["Source", "Limitation"])
    for row in warnings:
        append_safe(coverage, row)
    dictionary = book.create_sheet("Column guide")
    for row in [["Column", "Meaning"], ["price", "Raw item price; vendor modifiers are not applied"],
                ["weight", "Raw weight parameter, usually kg; ammo uses per-round weight"],
                ["flatweight", "Fixed weight added by the item's weight implementation"],
                ["baseweight / varweight", "Artifact weight parameters; runtime rolls are not resolved"],
                ["FireDelay", "Seconds between shots; theoretical RPM = 60 / FireDelay"],
                ["hp / hpvar", "Base mutant health and random variation"],
                ["BR / FBR / FBRAP", "Raw mutant resistance parameters, as named in Lua"],
                ["br / fbr / ar / far", "Raw item protection parameters, as named in Lua"],
                ["interval", "Anomaly-controller spacing parameter in Source units"],
                ["*Artifacts", "JSON array of item IDs; order and duplicates are preserved"],
                ["Other columns", "Exact Lua attribute names; consult the linked source before interpreting effects"]]:
        append_safe(dictionary, row)
    for sheet in book:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(65, max(15, max(len(str(c.value or "")) for c in col) + 2))
    book.save(output / "balance.xlsx")
    (output / "coverage.json").write_text(json.dumps(warnings, indent=2) + "\n")
    print(f"Exported {len(entities)} definitions; {len(warnings)} coverage notices to {output}")


def trusted_baseline(root, manifest):
    """Reconstruct baseline from an ancestor Git commit; never trust workbook paths."""
    commit = manifest.get("commit", "")
    if manifest.get("schema") != SCHEMA:
        raise BalanceError("Unsupported baseline schema; generate a fresh workbook and baseline.json.")
    if not re.fullmatch(r"[0-9a-f]{40,64}", commit):
        raise BalanceError("Invalid baseline commit")
    try:
        git(root, "merge-base", "--is-ancestor", commit, "HEAD")
        archive = subprocess.check_output(["git", "-C", str(root), "archive", commit, "gamemodes/ix_cotz", "lua"])
    except subprocess.CalledProcessError as exc:
        raise BalanceError("Baseline commit must be available in Git history and an ancestor of HEAD (fetch full history).") from exc
    with tempfile.TemporaryDirectory(prefix="cotz-baseline-") as folder:
        base = Path(folder)
        with tarfile.open(fileobj=io.BytesIO(archive)) as tar:
            for entry in tar:
                if not entry.isfile() or not entry.name.endswith(".lua"):
                    continue
                target = base / entry.name
                if not target.resolve().is_relative_to(base.resolve()):
                    raise BalanceError("Invalid archive path")
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(tar.extractfile(entry).read())
        entities, _ = scan(base)
    expected = {"schema": SCHEMA, "commit": commit, "entities": entities}
    if manifest != expected:
        raise BalanceError("Baseline does not match committed Lua. Export committed source and use its original baseline.json.")


def validate(value, original, field):
    if isinstance(original, list):
        try:
            parsed = json.loads(value)
        except (ValueError, TypeError) as exc:
            raise BalanceError("Expected a JSON array of item IDs") from exc
        if not isinstance(parsed, list) or any(not isinstance(v, str) or not re.fullmatch(r"[_a-zA-Z0-9]+", v) for v in parsed):
            raise BalanceError("Expected a JSON array of item IDs")
        return parsed
    if isinstance(original, bool):
        if not isinstance(value, bool):
            raise BalanceError("Expected TRUE or FALSE")
        return value
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise BalanceError("Expected a finite number")
    if value < 0 and not (field in {"Primary.ClipSize", "Primary.DefaultClip"} and value == -1):
        raise BalanceError("Negative values are not supported")
    if field in INTEGER_FIELDS and value != int(value):
        raise BalanceError("Expected an integer")
    if field in {"width", "height", "FireDelay", "interval"} and value == 0:
        raise BalanceError("Expected a positive value")
    return value


def lua_value(value):
    if isinstance(value, list):
        return "{" + ", ".join(json.dumps(v) for v in value) + "}"
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(int(value)) if value == int(value) else repr(value)


def import_book(root, workbook, baseline, report, apply=False):
    manifest = json.loads(baseline.read_text())
    trusted_baseline(root, manifest)
    book = load_workbook(workbook, data_only=False)
    if "Metadata" not in book:
        raise BalanceError("Missing Metadata sheet")
    metadata = {row[0]: row[1] for row in book["Metadata"].iter_rows(values_only=True) if len(row) >= 2}
    if metadata.get("Baseline SHA256") != digest(manifest) or metadata.get("Commit") != manifest["commit"] or metadata.get("Schema") != SCHEMA:
        raise BalanceError("Workbook and baseline do not match")
    current, _ = scan(root)
    key = lambda e: (e["tab"], e["id"], e["source"])
    current = {key(e): e for e in current}
    expected = {key(e): e for e in manifest["entities"]}
    changes = []
    replacements = defaultdict(list)
    errors = []
    seen = set()
    item_ids = {e["id"] for e in current.values() if e["tab"] == "Items"}
    for tab, (rows, columns) in workbook_layout(manifest["entities"]).items():
        if tab not in book:
            errors.append(f"Missing sheet: {tab}")
            continue
        sheet = book[tab]
        headers = [cell.value for cell in sheet[1]]
        if headers != IDENTITY + columns:
            errors.append(f"{tab}: columns have been changed")
            continue
        sheet_expected = {(e["id"], e["source"]): e for e in rows}
        for row in sheet.iter_rows(min_row=2):
            values = [c.value for c in row]
            if all(v is None for v in values):
                continue
            old = sheet_expected.get((str(values[0]), values[1]))
            ident = key(old) if old else None
            label = f"{tab}/{values[0]} ({values[1]})"
            if old is None or ident in seen:
                errors.append(f"{label}: unknown or duplicate row")
                continue
            seen.add(ident)
            identity = [old["id"], old["source"], old["name"], old["base"]]
            # XLSX readers represent saved empty strings as empty cells.
            identity = [None if v == "" else v for v in identity]
            if values[:len(IDENTITY)] != identity:
                errors.append(f"{label}: read-only identity/metadata changed")
                continue
            for index, field in enumerate(headers[len(IDENTITY):], len(IDENTITY)):
                cell = row[index]
                value = cell.value
                original = old["fields"].get(field)
                if cell.data_type == "f":
                    errors.append(f"{label}/{field}: formulas are not accepted")
                    continue
                if original is None:
                    if value is not None:
                        errors.append(f"{label}/{field}: absent or unsupported field is read-only")
                    continue
                # XLSX may deserialize an integral float as int. Booleans must
                # remain distinct even though Python considers True == 1.
                if value == display(original["value"]) and isinstance(value, bool) == isinstance(original["value"], bool):
                    continue
                try:
                    proposed = validate(value, original["value"], field)
                    if proposed == original["value"]:
                        continue
                    now = current.get(ident, {}).get("fields", {}).get(field)
                    if now is None:
                        raise BalanceError("Source field removed or no longer unambiguous")
                    if now["value"] == proposed:
                        continue
                    if now["value"] != original["value"]:
                        raise BalanceError(f"Conflict: source is now {now['value']!r}")
                    if isinstance(proposed, list) and any(v not in item_ids for v in proposed):
                        raise BalanceError("Artifact pool references an unknown item ID")
                    path = root / old["source"]
                    if path.is_symlink() or not path.resolve().is_relative_to(root.resolve()):
                        raise BalanceError("Source path escapes repository or is a symlink")
                    replacements[path].append((now["start"], now["end"], lua_value(proposed)))
                    changes.append({"source": old["source"], "id": old["id"], "field": field,
                                    "before": now["value"], "after": proposed})
                except BalanceError as exc:
                    errors.append(f"{label}/{field}: {exc}")
    missing = set(expected) - seen
    if missing:
        errors.append(f"{len(missing)} baseline rows missing; deleting definitions is not supported")
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps({"baseline_commit": manifest["commit"], "applied": apply and not errors,
                                 "changes": changes, "errors": errors}, indent=2) + "\n")
    if errors:
        raise BalanceError("Import rejected; no Lua files changed:\n" + "\n".join(errors))
    prepared = {}
    for path, edits in replacements.items():
        text = path.read_bytes().decode("utf-8")
        for start, end, replacement in sorted(edits, reverse=True):
            text = text[:start] + replacement + text[end:]
        prepared[path] = text.encode("utf-8")
    if apply:
        for path, data in prepared.items():
            path.write_bytes(data)
    print(f"{'Applied' if apply else 'Dry run:'} {len(changes)} changes in {len(prepared)} files. Report: {report}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    commands = parser.add_subparsers(dest="command", required=True)
    exp = commands.add_parser("export")
    exp.add_argument("--output", type=Path, default=Path("docs/balance/output"))
    exp.add_argument("--preview", action="store_true", help="Allow uncommitted source for inspection; cannot import until the baseline matches a commit")
    imp = commands.add_parser("import")
    imp.add_argument("--workbook", type=Path, required=True)
    imp.add_argument("--baseline", type=Path, required=True)
    imp.add_argument("--report", type=Path, default=Path("docs/balance/output/import-report.json"))
    imp.add_argument("--apply", action="store_true", help="Write files after validating the entire workbook")
    args = parser.parse_args()
    try:
        if args.command == "export":
            export(args.root.resolve(), args.output, args.preview)
        else:
            import_book(args.root.resolve(), args.workbook, args.baseline, args.report, args.apply)
    except (BalanceError, OSError, ValueError, subprocess.CalledProcessError) as exc:
        print(str(exc), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
