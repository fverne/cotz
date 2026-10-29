import contextlib
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import balance


class ExtractionTests(unittest.TestCase):
    def test_glua_comments_strings_and_function_scopes(self):
        source = '''\ufeff-- ITEM.price = 900
--[=[ ITEM.price = 800 ]=]
/* ITEM.price = 700 */
local text = "ITEM.price = 600"
ITEM.price = 365 // comment
ITEM.weight = 0.050
function ITEM:Example()
    for i = 1, 2 do
        if i == 1 then ITEM.price = 99 end
    end
end
if SERVER then ITEM.width = 3 end
ITEM.height = 1
'''
        fields, issues = balance.extract(source, 'ITEM', balance.ITEM_FIELDS)
        self.assertEqual(set(fields), {'price', 'weight', 'height'})
        self.assertEqual(fields['price']['value'], 365)
        self.assertEqual(source[fields['weight']['start']:fields['weight']['end']], '0.050')
        self.assertEqual(issues, [])

    def test_repeated_and_computed_values_are_not_editable(self):
        source = '''ITEM.price = 10
ITEM.price = GetPrice()
ITEM.weight = 1 + 2
ITEM.height = 1
 * scale
ITEM.width = 2
'''
        fields, issues = balance.extract(source, 'ITEM', balance.ITEM_FIELDS)
        self.assertEqual(set(fields), {'width'})
        self.assertEqual(len(issues), 3)

    def test_table_fields_and_signed_exponents(self):
        fields, _ = balance.extract('SWEP.Primary.ClipSize = 30\nSWEP.FireDelay = 9.2e-2\n', 'SWEP', balance.WEAPON_FIELDS)
        self.assertEqual(fields['Primary.ClipSize']['value'], 30)
        self.assertEqual(fields['FireDelay']['value'], .092)
        fields, _ = balance.extract('ITEM.weight = -0.5\n', 'ITEM', balance.ITEM_FIELDS)
        self.assertEqual(fields['weight']['value'], -.5)

    def test_parent_table_overwrite_disables_earlier_field(self):
        fields, issues = balance.extract('SWEP.Primary.ClipSize = 30\nSWEP.Primary = {}\n', 'SWEP', balance.WEAPON_FIELDS)
        self.assertEqual(fields, {})
        self.assertTrue(any('overwritten' in issue for issue in issues))

    def test_unbalanced_source_is_not_editable(self):
        with self.assertRaises(balance.BalanceError):
            balance.extract('ITEM.price = 365\nfunction ITEM:Broken()\n', 'ITEM', balance.ITEM_FIELDS)


class RoundTripTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='cotz-balance-test-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.item = self.root / 'gamemodes/ix_cotz/schema/items/aid/sh_medkit.lua'
        self.original = b'\xef\xbb\xbfITEM.name = "Medkit"\r\nITEM.price = 365 // preserve this\r\nITEM.weight = 0.050\r\n'
        self.write(self.item, self.original)
        self.write(self.item.with_name('sh_other.lua'), b'ITEM.name = "Other"\nITEM.price = 100\n')
        # Excluded by Helix's non-recursive item loader.
        self.write(self.item.parent / 'old/sh_inactive.lua', b'ITEM.price = 999\n')
        self.write(self.root / 'lua/weapons/test/shared.lua', b'SWEP.Damage = 78\nSWEP.Primary.ClipSize = 30\n')
        self.anomaly = self.root / 'gamemodes/ix_cotz/plugins/anomalycontroller/sh_anomalydefs.lua'
        self.write(self.anomaly, b'PLUGIN.anomalydefs[1] = {\n name = "Burner",\n entityname = "anom_jarka",\n interval = 128,\n commonArtifacts = {"medkit"},\n}\n')
        self.run_git('init', '-q')
        self.run_git('config', 'user.name', 'Balance tests')
        self.run_git('config', 'user.email', 'balance@example.invalid')
        self.commit()
        self.out = self.root / 'output'
        with contextlib.redirect_stdout(io.StringIO()):
            balance.export(self.root, self.out)
        self.workbook = self.out / 'balance.xlsx'
        self.baseline = self.out / 'baseline.json'
        self.report = self.out / 'report.json'

    def write(self, path, data):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)

    def run_git(self, *args):
        return subprocess.check_output(['git', '-C', str(self.root), *args], stderr=subprocess.PIPE)

    def commit(self):
        self.run_git('add', 'gamemodes', 'lua')
        self.run_git('commit', '-qm', 'Fixture')

    def edit(self, field, value, tab='Items - aid', ident='medkit'):
        book = load_workbook(self.workbook)
        sheet = book[tab]
        headers = [cell.value for cell in sheet[1]]
        for row in sheet.iter_rows(min_row=2):
            if str(row[0].value) == ident:
                row[headers.index(field)].value = value
                break
        else:
            self.fail('Fixture row missing')
        book.save(self.workbook)
        book.close()

    def apply(self, write=True):
        with contextlib.redirect_stdout(io.StringIO()):
            balance.import_book(self.root, self.workbook, self.baseline, self.report, write)
        return json.loads(self.report.read_text())

    def test_unchanged_roundtrip_is_byte_identical(self):
        self.assertEqual(self.apply()['changes'], [])
        self.assertEqual(self.item.read_bytes(), self.original)
        data = json.loads(self.baseline.read_text())
        self.assertNotIn('inactive', {e['id'] for e in data['entities']})

    def test_progression_tabs_are_read_only_views_and_do_not_change_import(self):
        self.write(self.root / 'gamemodes/ix_cotz/plugins/progression/definitions/sh_test.lua', b'''
ix.progression.Register("test", {
name = "Test", keyNpc = "Guide", progMap = "map", defaultActive = true,
GetItemIds = function() return {["medkit"] = 5} end,
})
''')
        self.commit()
        with contextlib.redirect_stdout(io.StringIO()):
            balance.export(self.root, self.out)
        book = load_workbook(self.workbook)
        self.assertEqual(book['Progression steps'].max_row, 2)
        sheet = book['Progression hand-ins']
        self.assertEqual(sheet.cell(2, 7).value, 5)
        self.assertEqual(sheet.cell(2, 7).fill.fgColor.rgb, '00F2F2F2')
        self.assertTrue(sheet.cell(2, 7).protection.locked)
        sheet.cell(2, 7, 1)
        book.save(self.workbook)
        book.close()
        self.edit('price', 400)
        self.assertEqual(len(self.apply()['changes']), 1)
        self.assertEqual(self.item.read_bytes(), self.original.replace(b'365', b'400'))

    def test_single_edit_preserves_bom_crlf_comments_and_other_literals(self):
        self.edit('price', 400)
        self.assertEqual(len(self.apply()['changes']), 1)
        self.assertEqual(self.item.read_bytes(), self.original.replace(b'365', b'400'))

    def test_dry_run_does_not_write(self):
        self.edit('price', 400)
        report = self.apply(False)
        self.assertEqual(len(report['changes']), 1)
        self.assertFalse(report['applied'])
        self.assertEqual(self.item.read_bytes(), self.original)

    def test_unchanged_negative_sentinel_and_integer_float_roundtrip(self):
        path = self.root / 'lua/weapons/test/shared.lua'
        path.write_bytes(b'SWEP.Damage = 78\nSWEP.Primary.ClipSize = -1\n')
        self.commit()
        with contextlib.redirect_stdout(io.StringIO()):
            balance.export(self.root, self.out)
        self.assertEqual(self.apply()['changes'], [])

    def test_unrelated_source_change_is_preserved(self):
        self.edit('price', 400)
        modified = self.original.replace(b'0.050', b'0.075') + b'-- concurrent comment\r\n'
        self.item.write_bytes(modified)
        self.apply()
        self.assertEqual(self.item.read_bytes(), modified.replace(b'365', b'400'))

    def test_conflict_rejects_entire_import(self):
        self.edit('price', 400)
        self.edit('price', 200, ident='other')
        modified = self.original.replace(b'365', b'450')
        self.item.write_bytes(modified)
        with self.assertRaisesRegex(balance.BalanceError, 'Conflict'):
            self.apply()
        self.assertEqual(self.item.read_bytes(), modified)
        self.assertIn(b'100', self.item.with_name('sh_other.lua').read_bytes())

    def test_source_already_matches_proposal(self):
        self.edit('price', 400)
        self.item.write_bytes(self.original.replace(b'365', b'400'))
        self.assertEqual(self.apply()['changes'], [])

    def test_reordered_rows_are_safe(self):
        self.edit('price', 400)
        book = load_workbook(self.workbook)
        sheet = book['Items - aid']
        rows = list(sheet.values)[1:]
        sheet.delete_rows(2, sheet.max_row)
        for row in reversed(rows):
            sheet.append(row)
        book.save(self.workbook)
        book.close()
        self.apply()
        self.assertEqual(self.item.read_bytes(), self.original.replace(b'365', b'400'))

    def test_formulas_blanks_negatives_and_wrong_types_are_rejected(self):
        for value in ['=1+1', None, -1, True, '400']:
            with self.subTest(value=value):
                self.edit('price', value)
                with self.assertRaises(balance.BalanceError):
                    self.apply()
                self.assertEqual(self.item.read_bytes(), self.original)

    def test_missing_duplicate_or_unknown_rows_are_rejected(self):
        for mode in ['missing', 'duplicate', 'unknown']:
            with self.subTest(mode=mode):
                with contextlib.redirect_stdout(io.StringIO()):
                    balance.export(self.root, self.out)
                book = load_workbook(self.workbook)
                sheet = book['Items - aid']
                if mode == 'missing':
                    sheet.delete_rows(2)
                elif mode == 'duplicate':
                    sheet.append([cell.value for cell in sheet[2]])
                else:
                    sheet.cell(2, 1, '../../escape')
                book.save(self.workbook)
                book.close()
                with self.assertRaises(balance.BalanceError):
                    self.apply()

    def test_tampered_manifest_is_rejected_even_with_matching_workbook_hash(self):
        data = json.loads(self.baseline.read_text())
        data['entities'][0]['source'] = '../../escape.lua'
        self.baseline.write_text(json.dumps(data))
        book = load_workbook(self.workbook)
        book['Metadata'].cell(3, 2, balance.digest(data))
        book.save(self.workbook)
        book.close()
        with self.assertRaisesRegex(balance.BalanceError, 'Baseline does not match'):
            self.apply()

    def test_artifact_pool_edits_and_unknown_references(self):
        self.edit('commonArtifacts', '["other", "medkit"]', 'Anomaly fields', '1')
        self.apply()
        self.assertIn(b'{"other", "medkit"}', self.anomaly.read_bytes())
        self.edit('commonArtifacts', '["missing_item"]', 'Anomaly fields', '1')
        # Reset current source to avoid a three-way conflict masking reference validation.
        self.run_git('checkout', '--', str(self.anomaly.relative_to(self.root)))
        with self.assertRaisesRegex(balance.BalanceError, 'unknown item ID'):
            self.apply()

    def test_export_requires_committed_configuration(self):
        self.item.write_bytes(self.original.replace(b'365', b'400'))
        with self.assertRaisesRegex(balance.BalanceError, 'Baseline does not match'):
            balance.export(self.root, self.out)

    def test_duplicate_item_ids_are_read_only(self):
        self.write(self.root / 'gamemodes/ix_cotz/plugins/example/items/aid/sh_medkit.lua', b'ITEM.price = 500\n')
        entities, warnings = balance.scan(self.root)
        self.assertEqual([e['fields'] for e in entities if e['id'] == 'medkit'], [{}, {}])
        self.assertTrue(any('Duplicate ID' in message for _, message in warnings))


    def test_category_sheets_use_inherited_names_and_only_relevant_fields(self):
        self.write(self.root / 'gamemodes/ix_cotz/schema/items/base/sh_aid.lua', b'ITEM.category = "Medical"\n')
        self.write(self.root / 'gamemodes/ix_cotz/schema/items/newammo/sh_ammo.lua', b'ITEM.category = "Ammunition"\nITEM.ammoAmount = 60\nITEM.weight = 0.01\n')
        self.write(self.root / 'gamemodes/ix_cotz/schema/items/aid/sh_special.lua', b'ITEM.category = "Special Aid"\nITEM.restore = 5\n')
        self.commit()
        with contextlib.redirect_stdout(io.StringIO()):
            balance.export(self.root, self.out)
        book = load_workbook(self.workbook)
        self.assertNotIn('Items', book.sheetnames)
        medical = [cell.value for cell in book['Items - Medical'][1]]
        ammo = [cell.value for cell in book['Items - Ammunition'][1]]
        special = [cell.value for cell in book['Items - Special Aid'][1]]
        self.assertEqual(medical, ['ID', 'Source', 'Name', 'Base', 'price', 'weight'])
        self.assertEqual(ammo, ['ID', 'Source', 'Name', 'Base', 'ammoAmount', 'weight'])
        self.assertEqual(special, ['ID', 'Source', 'Name', 'Base', 'restore'])
        for sheet in book:
            headers = [cell.value for cell in sheet[1]]
            self.assertNotIn('Availability', headers)
            self.assertNotIn('Class / Ammo', headers)
        book.close()
        self.edit('price', 400, tab='Items - Medical')
        self.edit('ammoAmount', 80, tab='Items - Ammunition', ident='ammo')
        self.assertEqual(len(self.apply()['changes']), 2)
        self.assertEqual(self.item.read_bytes(), self.original.replace(b'365', b'400'))

    def test_row_moved_to_another_category_is_rejected(self):
        self.write(self.root / 'gamemodes/ix_cotz/schema/items/misc/sh_misc.lua', b'ITEM.price = 10\nITEM.weight = 1\n')
        self.commit()
        with contextlib.redirect_stdout(io.StringIO()):
            balance.export(self.root, self.out)
        book = load_workbook(self.workbook)
        row = [cell.value for cell in book['Items - aid'][2]]
        book['Items - aid'].delete_rows(2)
        book['Items - misc'].append(row)
        book.save(self.workbook)
        book.close()
        with self.assertRaisesRegex(balance.BalanceError, 'unknown or duplicate row'):
            self.apply()

    def test_legacy_schema_requires_fresh_export(self):
        manifest = json.loads(self.baseline.read_text())
        manifest['schema'] = 1
        self.baseline.write_text(json.dumps(manifest))
        with self.assertRaisesRegex(balance.BalanceError, 'fresh workbook'):
            self.apply()


class LayoutTests(unittest.TestCase):
    def test_category_names_are_valid_and_unique_excel_titles(self):
        categories = ['Medical', 'medical', 'A/B', 'A:B', 'x' * 60]
        entities = [{'tab': 'Items', 'category': category, 'fields': {}} for category in categories]
        names = list(balance.workbook_layout(entities))[:len(categories)]
        self.assertEqual(len(set(name.casefold() for name in names)), len(categories))
        for name in names:
            self.assertLessEqual(len(name), 31)
            self.assertFalse(any(c in name for c in '\\/*?:[]'))


if __name__ == '__main__':
    unittest.main()
